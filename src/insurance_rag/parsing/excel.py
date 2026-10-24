"""Read-only worksheet adapter; formulas and cached values remain separate."""
from datetime import date,datetime
from decimal import Decimal
import importlib.metadata
from .common import stable_hash


def excel_versions():
    result={}
    for name in ('openpyxl','xlrd'):
        try:result[name]=importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:pass
    return result


def text(value):
    if value is None:return ''
    if isinstance(value,(date,datetime)):return value.isoformat()
    if isinstance(value,bool):return 'TRUE' if value else 'FALSE'
    if isinstance(value,(int,float,Decimal)):return format(Decimal(str(value)),'f')
    return str(value)


def extract_workbook(file,config,parse_key,max_cells=250000):
    from pathlib import Path
    kind=Path(file).suffix.lower()
    sheets=[]
    if kind=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(file,data_only=False,keep_links=False)
        cached=load_workbook(file,data_only=True,keep_links=False)
        try:
            for ws in book.worksheets:
                nrows,ncols=ws.max_row,ws.max_column
                if nrows*ncols>max_cells:raise ValueError('worksheet_cell_limit_exceeded')
                merged=[(m.min_row-1,m.max_row,m.min_col-1,m.max_col) for m in ws.merged_cells.ranges]
                values=[]
                for row in ws.iter_rows():
                    for cell in row:
                        formula=cell.value if cell.data_type=='f' else None
                        value=cached[ws.title].cell(cell.row,cell.column).value if formula else cell.value
                        values.append({'r':cell.row-1,'c':cell.column-1,'value':value,'formula':formula,
                            'data_type':cell.data_type,'number_format':cell.number_format,
                            'address':cell.coordinate,'hidden_row':bool(ws.row_dimensions[cell.row].hidden),
                            'hidden_column':bool(ws.column_dimensions[cell.column_letter].hidden) if hasattr(cell,'column_letter') else False})
                sheets.append({'name':ws.title,'rows':nrows,'cols':ncols,'merged':merged,'values':values,
                               'hidden':ws.sheet_state!='visible','legacy':False})
        finally:book.close();cached.close()
    elif kind=='.xls':
        import xlrd
        book=xlrd.open_workbook(str(file),formatting_info=True,on_demand=True)
        try:
            for ws in book.sheets():
                if ws.nrows*ws.ncols>max_cells:raise ValueError('worksheet_cell_limit_exceeded')
                values=[]
                for r in range(ws.nrows):
                    for c in range(ws.ncols):
                        cell=ws.cell(r,c);value=cell.value
                        if cell.ctype==xlrd.XL_CELL_DATE:value=datetime(*xlrd.xldate_as_tuple(value,book.datemode))
                        if cell.ctype in (xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK):value=None
                        values.append({'r':r,'c':c,'value':value,'formula':None,'data_type':str(cell.ctype),
                            'number_format':book.format_map[book.xf_list[cell.xf_index].format_key].format_str,
                            'address':xlrd.cellname(r,c),'hidden_row':bool(ws.rowinfo_map.get(r) and ws.rowinfo_map[r].hidden),
                            'hidden_column':bool(ws.colinfo_map.get(c) and ws.colinfo_map[c].hidden)})
                sheets.append({'name':ws.name,'rows':ws.nrows,'cols':ws.ncols,'merged':ws.merged_cells,
                               'values':values,'hidden':bool(ws.visibility),'legacy':True})
        finally:book.release_resources()
    else:raise ValueError('unsupported_workbook_format')
    pages=[];tables=[];blocks=[]
    for number,sheet in enumerate(sheets,1):
        reasons=['worksheet_semantics_not_verified']
        if sheet['hidden']:reasons.append('hidden_worksheet')
        if sheet['legacy']:reasons.append('legacy_formula_expressions_unavailable')
        occupied={};anchors={}
        for r0,r1,c0,c1 in sheet['merged']:
            anchors[(r0,c0)]=(r1-r0,c1-c0)
            for r in range(r0,r1):
                for c in range(c0,c1):occupied[(r,c)]=(r0,c0)
        tid=stable_hash(['worksheet-table-v1',parse_key,number]);cells=[]
        for original in sheet['values']:
            r,c=original['r'],original['c']
            if (r,c) in occupied and occupied[(r,c)]!=(r,c):continue
            rs,cs=anchors.get((r,c),(1,1));formula=original['formula'];value=original['value']
            raw=text(value)
            formula_reasons=[]
            if formula:
                formula_reasons=['formula_cache_missing' if value is None else 'formula_cache_not_verified']
                reasons.extend(formula_reasons)
                if '[' in formula:reasons.append('external_formula_reference')
            if original['hidden_row'] or original['hidden_column']:reasons.append('hidden_cells')
            box=[c,r,c+cs,r+rs]
            cells.append({'cell_id':stable_hash([tid,r,c]),'row_index':r,'column_index':c,
                'row_span':rs,'column_span':cs,'raw_text':raw,'normalized_text':raw,
                'bounding_box':box,'native_bounding_box':box,'numeric_tokens':[],
                'extraction_confidence':None,'source_block_ids':[],'header_candidate':r==0,
                'formula':formula,'cached_value':text(value) if formula and value is not None else None,
                'value_type':original['data_type'],'number_format':original['number_format'],
                'review_reasons':formula_reasons,'original_reference':{'page_number':number,'sheet_name':sheet['name'],
                'cell_address':original['address'],'coordinate_system':'worksheet_grid','hidden_row':original['hidden_row'],
                'hidden_column':original['hidden_column']}})
        width,height=max(1,sheet['cols']),max(1,sheet['rows'])
        pages.append({'page_number':number,'sheet_name':sheet['name'],'width':width,'height':height,
            'rotation':0,'coordinate_system':'worksheet_grid','status':'completed_with_warnings',
            'content_type':'worksheet','review_reasons':sorted(set(reasons)),'error':None})
        if not sheet['rows'] or not sheet['cols']:continue
        table={'table_id':tid,'page_number':number,'sheet_name':sheet['name'],'row_count':sheet['rows'],
            'column_count':sheet['cols'],'bounding_box':[0,0,width,height],'native_bounding_box':[0,0,width,height],
            'cells':cells,'extraction_method':'xlrd_cached_cells' if sheet['legacy'] else 'openpyxl_worksheet_cells',
            'header_rows_candidate':[0],'review_status':'pending_review','review_reasons':sorted(set(reasons)),
            'raw_text':'\n'.join(c['raw_text'] for c in cells),'continuation_of':None}
        tables.append(table)
        blocks.append({'block_id':stable_hash(['worksheet-block-v1',tid]),'block_type':'table',
            'raw_text':table['raw_text'],'normalized_text':table['raw_text'],'page_number':number,
            'reading_order':0,'bounding_box':table['bounding_box'],'native_bounding_box':table['native_bounding_box'],
            'parent_section_id':None,'clause_id':None,'table_id':tid,'extraction_method':table['extraction_method'],
            'extraction_confidence':None,'cleaning_changes':[],
            'original_reference':{'page_number':number,'sheet_name':sheet['name'],'coordinate_system':'worksheet_grid'}})
    return {'page_count':len(pages),'pages':pages,'blocks':blocks,'tables':tables,'sections':[],'errors':[]}
